"""Downloads — an answer's data as CSV or Excel, and a report as one HTML file.

The browser sends what it already shows: a result's columns and rows (already
governed — masked and denied columns were applied when the query ran), and
for a report its sections, each with text, SQL, a table and a chart image.
Nothing is queried here, so a download can never show more than its reader
saw. Building the files server-side keeps three things in one place:

- SPREADSHEET SAFETY. A cell starting with = + - @ is a formula to Excel. In
  a CSV such a text value is prefixed with ' (OWASP's advice for CSV
  injection); in an .xlsx every text cell is written as text, never as a
  formula, so =HYPERLINK(…) arrives as the words it was.
- ESCAPING. The report is static HTML: every value is escaped, there is no
  script, and a chart is embedded only as a PNG or SVG data-URL image (which
  a browser never runs).
- THE AUDIT. Every download is an activity row (action "export", its format
  in mode), like an emailed report.
"""
import base64
import csv
import html
import io
import re
import time
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel

from . import db
from .auth import current_user

router = APIRouter(prefix="/export", tags=["export"])

MAX_ROWS = 200_000          # rows per download (an answer holds far fewer)
MAX_COLUMNS = 1_000
MAX_SECTIONS = 60
REPORT_ROWS = 1_000         # table rows per report section
MAX_IMAGE_BYTES = 8 * 1024 * 1024
_EXCEL_CELL = 32_767        # Excel's limit on one cell's text
_IMAGE = re.compile(r"^data:image/(png|svg\+xml);base64,([A-Za-z0-9+/=\s]+)$")
_NUMBER = re.compile(r"^[-+]?(\d+\.?\d*|\.\d+)([eE][-+]?\d+)?$")
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class Table(BaseModel):
    title: str = "Studio export"
    columns: List[str]
    rows: List[list]
    sql: Optional[str] = None
    text: Optional[str] = None
    source: Optional[str] = None


class Section(BaseModel):
    title: str = ""
    text: Optional[str] = None
    image: Optional[str] = None          # data:image/png|svg+xml;base64,…
    columns: List[str] = []
    rows: List[list] = []
    sql: Optional[str] = None
    source: Optional[str] = None


class Report(BaseModel):
    title: str = "Studio report"
    subtitle: Optional[str] = None
    sections: List[Section]


# ── Helpers ──────────────────────────────────────────────────────────────

def filename(title, ext):
    """A safe file name from a title: words and dashes, at most 60 long."""
    stem = re.sub(r"[^\w-]+", "-", str(title or "")).strip("-_")[:60] or "studio-export"
    return f"{stem}.{ext}"


def _attachment(body, media_type, name):
    return Response(content=body, media_type=media_type,
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


def _bounded(columns, rows):
    if len(columns) > MAX_COLUMNS:
        raise HTTPException(413, f"At most {MAX_COLUMNS} columns can be downloaded")
    if len(rows) > MAX_ROWS:
        raise HTTPException(413, f"At most {MAX_ROWS:,} rows can be downloaded")


def csv_cell(value):
    """One value as CSV text; a text value that a spreadsheet would read as a
    formula gets a leading ' (a number, even a negative one, is left alone)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float)):
        return repr(value) if isinstance(value, float) else str(value)
    if isinstance(value, (dict, list)):
        import json
        value = json.dumps(value, ensure_ascii=False)
    text = str(value)
    if text.startswith(_FORMULA_START) and not _NUMBER.match(text):
        return "'" + text
    return text


def to_csv(columns, rows):
    """UTF-8 with a byte-order mark, so Excel reads accents and symbols."""
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\r\n")
    writer.writerow([csv_cell(c) for c in columns])
    for row in rows:
        writer.writerow([csv_cell(v) for v in row])
    return ("﻿" + out.getvalue()).encode("utf-8")


def to_xlsx(table, user_email):
    """The rows on a Data sheet (header bold and frozen, text as text), and
    what they are on an About sheet: title, question text, SQL, source."""
    import json

    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    widths = [min(50, max(8, len(str(c)) + 2)) for c in table.columns]
    for j, name in enumerate(table.columns, start=1):
        cell = ws.cell(row=1, column=j, value=str(name)[:_EXCEL_CELL])
        cell.data_type = "s"
        cell.font = Font(bold=True)
    for i, row in enumerate(table.rows, start=2):
        for j, value in enumerate(row[:len(table.columns)], start=1):
            if isinstance(value, (dict, list)):
                value = json.dumps(value, ensure_ascii=False)
            if isinstance(value, bool) or value is None or isinstance(value, (int, float)):
                ws.cell(row=i, column=j, value=value)
                continue
            text = str(value)[:_EXCEL_CELL]
            cell = ws.cell(row=i, column=j, value=text)
            cell.data_type = "s"                      # text, never a formula
            if j <= len(widths) and i <= 200:
                widths[j - 1] = min(50, max(widths[j - 1], len(text) + 2))
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    about = wb.create_sheet("About")
    facts = [("Title", table.title), ("Question / answer", table.text or ""),
             ("Source", table.source or ""), ("SQL", table.sql or ""),
             ("Rows", len(table.rows)),
             ("Exported", time.strftime("%Y-%m-%d %H:%M UTC", time.gmtime())),
             ("Exported by", user_email)]
    for i, (k, v) in enumerate(facts, start=1):
        about.cell(row=i, column=1, value=k).font = Font(bold=True)
        cell = about.cell(row=i, column=2, value=v if isinstance(v, int) else str(v)[:_EXCEL_CELL])
        if not isinstance(v, int):
            cell.data_type = "s"
    about.column_dimensions["A"].width = 18
    about.column_dimensions["B"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def checked_image(data_url):
    """A chart image the report may embed, or None: a PNG or SVG data URL of
    bounded size, re-encoded so nothing but base64 reaches the page."""
    if not data_url:
        return None
    m = _IMAGE.match(data_url.strip())
    if not m:
        raise HTTPException(400, "A chart image must be a PNG or SVG data URL")
    try:
        raw = base64.b64decode(re.sub(r"\s+", "", m.group(2)), validate=True)
    except ValueError:
        raise HTTPException(400, "A chart image is not valid base64")
    if len(raw) > MAX_IMAGE_BYTES:
        raise HTTPException(413, "A chart image is too large")
    return f"data:image/{m.group(1)};base64,{base64.b64encode(raw).decode()}"


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        import json
        value = json.dumps(value, ensure_ascii=False)
    return html.escape(str(value))


_STYLE = """
body{font:14px/1.5 -apple-system,"Segoe UI",Roboto,Helvetica,Arial,sans-serif;color:#1b1a19;
background:#ffffff;margin:0}
main{max-width:1000px;margin:0 auto;padding:32px 24px}
h1{font-size:24px;margin:0 0 4px}h2{font-size:17px;margin:0 0 8px}
.sub{color:#79766f;font-size:12px;margin-bottom:24px}
section{border-top:1px solid #e3e1de;padding:20px 0;break-inside:avoid-page}
.text{white-space:pre-wrap;margin:0 0 12px}.meta{color:#79766f;font-size:12px}
img{max-width:100%;height:auto;display:block;margin:8px 0 12px}
table{border-collapse:collapse;width:100%;font-size:12px;margin:8px 0}
th,td{border:1px solid #e3e1de;padding:4px 8px;text-align:left;vertical-align:top}
th{background:#f5f4f2;font-weight:600}
pre{background:#f5f4f2;border:1px solid #e3e1de;border-radius:6px;padding:10px;
white-space:pre-wrap;font-size:12px}
footer{color:#79766f;font-size:11px;border-top:1px solid #e3e1de;padding-top:12px}
@media print{main{padding:0}a{color:inherit}}
"""


def to_report_html(report, user_email):
    """One self-contained page: no script, no external request."""
    stamp = time.strftime("%Y-%m-%d %H:%M UTC", time.gmtime())
    parts = ['<!doctype html><html lang="en"><head><meta charset="utf-8">',
             '<meta name="viewport" content="width=device-width, initial-scale=1">',
             "<meta http-equiv=\"Content-Security-Policy\" content=\"default-src 'none'; "
             "img-src data:; style-src 'unsafe-inline'\">",
             f"<title>{html.escape(report.title)}</title><style>{_STYLE}</style></head>",
             f"<body><main><h1>{html.escape(report.title)}</h1>",
             f'<div class="sub">{html.escape(report.subtitle or "")}'
             f'{" · " if report.subtitle else ""}Exported {stamp}</div>']
    for s in report.sections:
        parts.append("<section>")
        if s.title:
            parts.append(f"<h2>{html.escape(s.title)}</h2>")
        if s.source:
            parts.append(f'<div class="meta">Source: {html.escape(s.source)}</div>')
        if s.text:
            parts.append(f'<p class="text">{html.escape(s.text)}</p>')
        image = checked_image(s.image)
        if image:
            parts.append(f'<img alt="{html.escape(s.title or "chart")}" src="{image}">')
        if s.columns and s.rows:
            shown = s.rows[:REPORT_ROWS]
            parts.append("<table><thead><tr>"
                         + "".join(f"<th>{_cell(c)}</th>" for c in s.columns)
                         + "</tr></thead><tbody>")
            for row in shown:
                parts.append("<tr>" + "".join(f"<td>{_cell(v)}</td>"
                                              for v in row[:len(s.columns)]) + "</tr>")
            parts.append("</tbody></table>")
            if len(s.rows) > len(shown):
                parts.append(f'<div class="meta">The first {len(shown):,} of {len(s.rows):,} '
                             "rows — download CSV or Excel for all of them.</div>")
        if s.sql:
            parts.append(f'<div class="meta">SQL</div><pre>{html.escape(s.sql)}</pre>')
        parts.append("</section>")
    parts.append(f"<footer>Exported by {html.escape(user_email)} from Studio.</footer>"
                 "</main></body></html>")
    return "".join(parts).encode("utf-8")


# ── Routes ───────────────────────────────────────────────────────────────

def _audit(user, fmt, title, sql=None, rows=0):
    db.log_activity(user, "export", prompt=str(title)[:300], sql=sql, row_count=rows, mode=fmt)


@router.post("/csv")
def export_csv(body: Table, user=Depends(current_user)):
    _bounded(body.columns, body.rows)
    _audit(user, "csv", body.title, body.sql, len(body.rows))
    return _attachment(to_csv(body.columns, body.rows), "text/csv; charset=utf-8",
                       filename(body.title, "csv"))


@router.post("/xlsx")
def export_xlsx(body: Table, user=Depends(current_user)):
    _bounded(body.columns, body.rows)
    data = to_xlsx(body, user["email"])
    _audit(user, "xlsx", body.title, body.sql, len(body.rows))
    return _attachment(data, XLSX, filename(body.title, "xlsx"))


@router.post("/report")
def export_report(body: Report, user=Depends(current_user)):
    if not body.sections:
        raise HTTPException(400, "A report needs at least one section")
    if len(body.sections) > MAX_SECTIONS:
        raise HTTPException(413, f"A report has at most {MAX_SECTIONS} sections")
    for s in body.sections:
        _bounded(s.columns, s.rows)
    page = to_report_html(body, user["email"])
    _audit(user, "report", body.title, rows=sum(len(s.rows) for s in body.sections))
    return _attachment(page, "text/html; charset=utf-8", filename(body.title, "html"))
